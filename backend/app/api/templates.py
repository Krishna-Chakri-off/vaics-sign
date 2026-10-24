"""Templates (FR-059..065), envelope-from-template, bulk send (FR-148..152)."""
import csv, io, json, re, secrets
from datetime import datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Form
from pydantic import BaseModel, EmailStr, Field as PField
from sqlalchemy import or_
from sqlalchemy.orm import Session
from ..core.config import settings
from ..core.db import get_db
from ..core.security import sha256_hex, random_token, numeric_code, keyed_hash
from ..models import Template, TemplateVersion, Envelope, EnvelopeDocument, Recipient, Field, BulkJob, Organization, User, OneTimeToken, FIELD_TYPES, now
from ..services import workflow as wf, audit, notify
from ..services.storage import storage
from .deps import Principal, perm, org_member, envelope_access, user_groups
from .serializers import envelope_out

router = APIRouter(tags=["templates & bulk send"])
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def t_out(t: Template, uid: int | None = None):
    return {**{k: getattr(t, k) for k in ("id", "name", "description", "sharing", "group_id", "version", "documents", "roles", "fields", "message", "settings", "owner_id", "archived", "created_at", "updated_at")},
            "favourite": uid in (t.favourite_of or [])}


def visible_templates(db, p: Principal):
    gids = user_groups(db, p)
    return db.query(Template).filter(Template.org_id == p.org.id, Template.archived.is_(False),
                                     or_(Template.owner_id == p.user.id, Template.sharing == "org", (Template.sharing == "group") & Template.group_id.in_(gids or [0])))


def tpl_access(db, p, tid, write=False) -> Template:
    t = db.get(Template, tid)
    if not t or t.org_id != p.org.id: raise HTTPException(404, "Template not found")
    if write:
        if t.owner_id != p.user.id and not p.can("templates.share"): raise HTTPException(403, "You cannot edit this template")
    elif t not in visible_templates(db, p).all():
        raise HTTPException(403, "You cannot use this template")
    return t


@router.get("/templates")
def list_templates(q: str | None = None, favourites: bool = False, p: Principal = Depends(org_member), db: Session = Depends(get_db)):
    rows = visible_templates(db, p).order_by(Template.updated_at.desc()).all()
    if q: rows = [t for t in rows if q.lower() in (t.name + " " + t.description).lower()]
    if favourites: rows = [t for t in rows if p.user.id in (t.favourite_of or [])]
    return [t_out(t, p.user.id) for t in rows]


class FromEnvelopeIn(BaseModel):
    envelope_id: int
    name: str
    description: str = ""
    sharing: str = PField("private", pattern="^(private|group|org)$")
    group_id: int | None = None


@router.post("/templates", status_code=201)
def create_template(body: FromEnvelopeIn, request: Request, p: Principal = Depends(perm("templates.create")), db: Session = Depends(get_db)):
    """Save an envelope (draft or sent) as a template: documents, roles (from recipients), fields, routing and message."""
    if body.sharing != "private" and not p.can("templates.share"): raise HTTPException(403, "Only template managers can share templates")
    env = envelope_access(db, p, body.envelope_id)
    ridx = {r.id: i for i, r in enumerate(env.recipients)}; didx = {d.id: i for i, d in enumerate(env.documents)}
    t = Template(org_id=p.org.id, owner_id=p.user.id, name=body.name, description=body.description, sharing=body.sharing, group_id=body.group_id, message=env.message,
                 documents=[{"name": d.name, "storage_key": d.storage_key, "sha256": d.sha256, "pages": d.pages, "text": d.text[:20000]} for d in env.documents],
                 roles=[{"role": r.role_name or f"Signer {i + 1}", "type": r.type, "order": r.order, "auth_method": r.auth_method, "name": "", "email": "", "locked": False,
                         "host_user_id": r.host_user_id} for i, r in enumerate(env.recipients)],
                 fields=[{"role_index": ridx[f.recipient_id], "doc_index": didx[f.document_id], **{k: getattr(f, k) for k in ("type", "page", "x", "y", "w", "h", "required", "read_only",
                          "reason_required", "label", "tooltip", "default_value", "options", "validation", "regex", "font_size", "group_name", "formula")}} for f in env.fields],
                 settings={"reminder_first_days": env.reminder_first_days, "reminder_every_days": env.reminder_every_days, "language": env.language})
    db.add(t); db.flush()
    db.add(TemplateVersion(template_id=t.id, version=1, snapshot=json.loads(json.dumps(t_out(t), default=str)), created_by=p.user.id))
    audit.admin_log(db, p.org.id, p.user, "template.created", t.name, after={"sharing": t.sharing}, req=request)
    db.commit()
    return t_out(t, p.user.id)


class TemplatePatch(BaseModel):
    name: str | None = None
    description: str | None = None
    sharing: str | None = PField(None, pattern="^(private|group|org)$")
    group_id: int | None = None
    roles: list[dict] | None = None
    fields: list[dict] | None = None
    message: str | None = None
    archived: bool | None = None


@router.put("/templates/{tid}")
def update_template(tid: int, body: TemplatePatch, request: Request, p: Principal = Depends(perm("templates.create")), db: Session = Depends(get_db)):
    t = tpl_access(db, p, tid, write=True)
    if body.sharing and body.sharing != "private" and not p.can("templates.share"): raise HTTPException(403, "Only template managers can share templates")
    for k, v in body.model_dump(exclude_none=True).items():
        setattr(t, k, v)
    t.version += 1
    db.add(TemplateVersion(template_id=t.id, version=t.version, snapshot=json.loads(json.dumps(t_out(t), default=str)), created_by=p.user.id))
    audit.admin_log(db, p.org.id, p.user, "template.updated", t.name, after={"version": t.version}, req=request)
    db.commit()
    return t_out(t, p.user.id)


@router.get("/templates/{tid}/versions")
def versions(tid: int, p: Principal = Depends(org_member), db: Session = Depends(get_db)):
    tpl_access(db, p, tid)
    return [{"version": v.version, "created_at": v.created_at, "created_by": v.created_by} for v in db.query(TemplateVersion).filter_by(template_id=tid).order_by(TemplateVersion.version)]


@router.post("/templates/{tid}/favourite")
def favourite(tid: int, p: Principal = Depends(org_member), db: Session = Depends(get_db)):
    t = tpl_access(db, p, tid)
    fav = list(t.favourite_of or [])
    t.favourite_of = [x for x in fav if x != p.user.id] if p.user.id in fav else fav + [p.user.id]
    db.commit()
    return {"favourite": p.user.id in t.favourite_of}


class RoleFill(BaseModel):
    name: str
    email: EmailStr
    phone: str | None = None


class UseIn(BaseModel):
    template_ids: list[int] = PField(min_length=1)  # several templates combine (FR-065)
    title: str | None = None
    message: str | None = None
    roles: list[RoleFill]
    prefill: dict[str, str] = {}  # field label -> value
    send: bool = False


def instantiate(db, org, user, templates: list[Template], title, message, roles: list[dict], prefill: dict, req=None, bulk_job_id=None, powerform_id=None) -> Envelope:
    base = templates[0]
    nroles = max(len(t.roles) for t in templates)
    if len(roles) < nroles: raise HTTPException(422, f"Provide name and email for all {nroles} roles")
    env = Envelope(org_id=org.id, owner_id=user.id, public_id=wf.public_id(), title=(title or base.name)[:200], message=message if message is not None else base.message,
                   template_id=base.id, template_version=base.version, bulk_job_id=bulk_job_id, powerform_id=powerform_id, language=base.settings.get("language", "en"),
                   reminder_first_days=base.settings.get("reminder_first_days", 2), reminder_every_days=base.settings.get("reminder_every_days", 2), is_sandbox=org.is_sandbox,
                   visibility=wf.org_settings(org)["default_visibility"])
    db.add(env); db.flush()
    recips = []
    for i, role in enumerate(base.roles):
        fill = roles[i]
        name = role.get("name") if role.get("locked") and role.get("name") else fill["name"]
        email = role.get("email") if role.get("locked") and role.get("email") else fill["email"]
        r = Recipient(envelope_id=env.id, role_name=role.get("role"), type=role.get("type", "signer"), name=name, email=email.lower(), phone=fill.get("phone"),
                      order=role.get("order", i + 1), auth_method=role.get("auth_method") or wf.org_settings(org)["default_auth"], host_user_id=role.get("host_user_id"))
        env.recipients.append(r); recips.append(r)
    db.flush()
    for t in templates:
        dmap = []
        for d in t.documents:
            nd = EnvelopeDocument(envelope_id=env.id, order=len(env.documents) + 1, name=d["name"], source_name=d["name"], storage_key=d["storage_key"], sha256=d["sha256"],
                                  source_sha256=d["sha256"], pages=d["pages"], text=d.get("text", ""), size=0)
            env.documents.append(nd); db.flush(); dmap.append(nd)
        for f in t.fields:
            if f["role_index"] >= len(recips): continue
            vals = {k: f.get(k) for k in ("type", "page", "x", "y", "w", "h", "required", "read_only", "reason_required", "label", "tooltip", "default_value", "options",
                                          "validation", "regex", "font_size", "group_name", "formula") if k in f}
            if f.get("label") and f["label"] in prefill: vals["default_value"] = prefill[f["label"]]
            env.fields.append(Field(envelope_id=env.id, document_id=dmap[f["doc_index"]].id, recipient_id=recips[f["role_index"]].id, **vals))
    env.version_hash = sha256_hex("|".join(d.sha256 for d in env.documents))
    wf.reindex(env)
    audit.record(db, env, "created", user.email, req, f"from template(s) {', '.join(f'{t.name} v{t.version}' for t in templates)}")
    return env


@router.post("/templates/use", status_code=201)
def use_template(body: UseIn, request: Request, p: Principal = Depends(perm("envelopes.send")), db: Session = Depends(get_db)):
    tpls = [tpl_access(db, p, t) for t in body.template_ids]
    env = instantiate(db, p.org, p.user, tpls, body.title, body.message, [r.model_dump() for r in body.roles], body.prefill, request)
    if body.send:
        db.flush(); wf.send(db, p.org, env, p.user, request)
    db.commit()
    return envelope_out(db, env, full=True)


# ---------------- bulk send ----------------
def parse_rows(data: bytes, filename: str) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise HTTPException(415, "Excel upload requires openpyxl; upload CSV instead")
        ws = openpyxl.load_workbook(io.BytesIO(data), read_only=True).active
        it = ws.iter_rows(values_only=True); head = [str(h or "").strip() for h in next(it)]
        return [{head[i]: ("" if v is None else str(v)).strip() for i, v in enumerate(row) if i < len(head)} for row in it if any(row)]
    text = data.decode("utf-8-sig", errors="replace")
    return [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(io.StringIO(text))]


@router.post("/bulk", status_code=201)
async def bulk_validate(request: Request, template_id: int = Form(...), name: str = Form("Bulk send"), file: UploadFile = File(...),
                        p: Principal = Depends(perm("bulk.send")), db: Session = Depends(get_db)):
    """Columns: for each template role N (1-based): roleN_name, roleN_email[, roleN_phone]; any other column = prefill by field label."""
    t = tpl_access(db, p, template_id)
    rows = parse_rows(await file.read(), file.filename or "list.csv")
    if not rows: raise HTTPException(422, "The list is empty")
    out, valid = [], 0
    for i, row in enumerate(rows, start=2):
        errs = []
        for n, role in enumerate(t.roles, start=1):
            if role.get("locked") and role.get("email"): continue
            nm, em = row.get(f"role{n}_name", ""), row.get(f"role{n}_email", "")
            if not nm: errs.append(f"role{n}_name is required")
            if not EMAIL_RE.match(em): errs.append(f"role{n}_email '{em}' is not a valid email")
        valid += not errs
        out.append({"row": i, "data": row, "errors": errs, "envelope_id": None})
    job = BulkJob(org_id=p.org.id, owner_id=p.user.id, template_id=t.id, name=name, total=len(rows), valid=valid, rows=out)
    db.add(job); audit.admin_log(db, p.org.id, p.user, "bulk.validated", name, after={"total": len(rows), "valid": valid}, req=request); db.commit()
    return bulk_out(job)


def bulk_out(j: BulkJob):
    return {k: getattr(j, k) for k in ("id", "name", "template_id", "status", "total", "valid", "sent", "rows", "created_at")}


@router.post("/bulk/{jid}/send")
def bulk_send(jid: int, request: Request, p: Principal = Depends(perm("bulk.send")), db: Session = Depends(get_db)):
    j = db.get(BulkJob, jid)
    if not j or j.org_id != p.org.id or j.owner_id != p.user.id: raise HTTPException(404, "Bulk job not found")
    if j.status != "validated": raise HTTPException(409, "Job already processed")
    t = tpl_access(db, p, j.template_id)
    j.status = "sending"; rows = [dict(r) for r in j.rows]; sent = 0
    for r in rows:
        if r["errors"]: continue
        try:
            with db.begin_nested():
                roles = [{"name": r["data"].get(f"role{n}_name", ""), "email": r["data"].get(f"role{n}_email", "x@invalid.invalid"), "phone": r["data"].get(f"role{n}_phone") or None}
                         for n in range(1, len(t.roles) + 1)]
                prefill = {k: v for k, v in r["data"].items() if not re.match(r"role\d+_", k)}
                env = instantiate(db, p.org, p.user, [t], f"{t.name} – {roles[0]['name']}", None, roles, prefill, request, bulk_job_id=j.id)
                db.flush(); wf.send(db, p.org, env, p.user, request)
                r["envelope_id"] = env.id; sent += 1
        except HTTPException as ex:
            r["errors"] = [str(ex.detail.get("message") if isinstance(ex.detail, dict) else ex.detail)]
    j.rows, j.sent, j.status = rows, sent, "done"
    audit.admin_log(db, p.org.id, p.user, "bulk.sent", j.name, after={"sent": sent}, req=request); db.commit()
    return bulk_out(j)


@router.get("/bulk")
def bulk_jobs(p: Principal = Depends(org_member), db: Session = Depends(get_db)):
    return [bulk_out(j) for j in db.query(BulkJob).filter_by(org_id=p.org.id, owner_id=p.user.id).order_by(BulkJob.id.desc())]


@router.get("/bulk/{jid}")
def bulk_job(jid: int, p: Principal = Depends(org_member), db: Session = Depends(get_db)):
    j = db.get(BulkJob, jid)
    if not j or j.org_id != p.org.id: raise HTTPException(404, "Bulk job not found")
    out = bulk_out(j)
    for r in out["rows"]:
        if r["envelope_id"]:
            e = db.get(Envelope, r["envelope_id"]); r["status"] = e.status if e else None
    return out


@router.get("/bulk/{jid}/report.csv")
def bulk_report(jid: int, p: Principal = Depends(org_member), db: Session = Depends(get_db)):
    from fastapi.responses import Response
    out = bulk_job(jid, p, db)
    b = io.StringIO(); w = csv.writer(b); w.writerow(["row", "envelope_id", "status", "errors"])
    for r in out["rows"]:
        e = db.get(Envelope, r["envelope_id"]) if r["envelope_id"] else None
        w.writerow([r["row"], e.public_id if e else "", r.get("status") or "not sent", "; ".join(r["errors"])])
    return Response(b.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="bulk-{jid}-report.csv"'})


class BulkActionIn(BaseModel):
    envelope_ids: list[int]
    action: str = PField(pattern="^(remind|void)$")
    reason: str = ""


@router.post("/envelopes-bulk-action")
def bulk_action(body: BulkActionIn, request: Request, p: Principal = Depends(perm("envelopes.send")), db: Session = Depends(get_db)):
    done, errors = 0, []
    for eid in body.envelope_ids[:1000]:
        try:
            env = envelope_access(db, p, eid, write=True)
            if body.action == "remind":
                done += sum(wf.remind(db, p.org, env, r, request, p.user.email) for r in wf.current_step(env)) > 0
            else:
                wf.void(db, p.org, env, p.user, body.reason or "Voided in bulk", request); done += 1
        except HTTPException as ex:
            errors.append({"envelope_id": eid, "error": str(ex.detail)})
    db.commit()
    return {"done": done, "errors": errors}
