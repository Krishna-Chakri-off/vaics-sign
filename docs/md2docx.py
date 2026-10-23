"""Tiny markdown -> docx converter for the brief (headings, paragraphs, bullets, numbered lists, tables, code blocks, **bold**)."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
IMG = None
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0D, 0x47, 0xA1)


def runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text.replace("<br/>", " ")):
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"
        elif part:
            p.add_run(part)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor); tcPr.append(s)


def convert(src, dst):
    doc = Document()
    for sec in doc.sections: sec.left_margin = sec.right_margin = Inches(0.8)
    import os
    logo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "../CMS-Frontend/public/vaics-logo-full.png")
    if os.path.exists(logo): doc.add_picture(logo, width=Inches(2))
    doc.styles["Normal"].font.name = "Segoe UI"; doc.styles["Normal"].font.size = Pt(10)
    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("```"):
            i += 1; code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            if l.strip() == "```mermaid" and IMG:
                doc.add_picture(IMG, width=Inches(6.5)); i += 1; continue
            p = doc.add_paragraph(); r = p.add_run("\n".join(code)); r.font.name = "Consolas"; r.font.size = Pt(7)
        elif l.strip() == "---":
            pass
        elif l.startswith("> "):
            p = doc.add_paragraph(); runs(p, l[2:]); p.paragraph_format.left_indent = Inches(0.3)
            for r in p.runs: r.font.color.rgb = RGBColor(0xE6, 0x00, 0x7E)
        elif l.startswith("#"):
            lvl = len(l) - len(l.lstrip("#")); h = doc.add_heading(l.lstrip("# ").strip(), level=min(lvl - 1, 3) if lvl > 1 else 0)
            for r in h.runs: r.font.color.rgb = NAVY
        elif l.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row[:len(rows[0])]):
                    cell = t.cell(ri, ci); cell.text = ""; runs(cell.paragraphs[0], val)
                    for r in cell.paragraphs[0].runs:
                        r.font.size = Pt(8.5)
                        if ri == 0: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
                    if ri == 0: shade(cell, "0D47A1")
            doc.add_paragraph()
        elif re.match(r"^\s*- ", l):
            runs(doc.add_paragraph(style="List Bullet"), l.split("- ", 1)[1])
        elif re.match(r"^\d+\. ", l):
            runs(doc.add_paragraph(style="List Number"), l.split(". ", 1)[1])
        elif l.strip():
            runs(doc.add_paragraph(), l)
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    IMG = sys.argv[3] if len(sys.argv) > 3 else None
    convert(sys.argv[1], sys.argv[2])
